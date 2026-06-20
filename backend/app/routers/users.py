import io

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from ..auth import get_current_user, hash_password, require_permission
from ..database import get_db
from ..models import Department, Report, User
from ..permissions import DEFAULT_ROLE_PERMISSIONS, Permission, Role
from ..schemas import ReportSummary, UserCreate, UserOut, UserUpdate
from ..services import (
    can_manage_user,
    can_view_report,
    descendant_department_ids,
)

router = APIRouter(prefix="/api/users", tags=["users"])

XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


def _to_out(db: Session, user: User) -> UserOut:
    out = UserOut.model_validate(user)
    out.department_name = user.department.name if user.department else None
    return out


def _require_can_manage(db: Session, current: User, target: User) -> None:
    if not can_manage_user(db, current, target):
        raise HTTPException(status_code=403, detail="没有权限管理该人员")


@router.get("", response_model=list[UserOut])
def list_users(
    department_id: int | None = None,
    role: str | None = None,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    is_admin = current.role == Role.ADMIN.value or (
        Permission.MANAGE_USERS.value in (current.permissions or [])
    )
    is_counselor = current.role == Role.COUNSELOR.value
    if not (is_admin or is_counselor):
        raise HTTPException(status_code=403, detail="没有权限查看人员")

    q = db.query(User)
    if not is_admin and is_counselor:
        # 咨询师仅可见所辖部门（含下级）的学员
        if not current.department_id:
            return []
        dept_ids = descendant_department_ids(db, current.department_id)
        q = q.filter(
            User.department_id.in_(dept_ids), User.role == Role.STUDENT.value
        )
    if department_id is not None:
        q = q.filter(User.department_id == department_id)
    if role:
        q = q.filter(User.role == role)
    return [_to_out(db, u) for u in q.all()]


@router.get("/staff", response_model=list[UserOut])
def list_staff(
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    """咨询师 + 管理员列表，供报告「咨询师签名」下拉选择。"""
    staff = (
        db.query(User)
        .filter(
            User.role.in_([Role.ADMIN.value, Role.COUNSELOR.value]),
            User.is_active == True,
        )
        .all()
    )
    return [_to_out(db, u) for u in staff]


@router.post("", response_model=UserOut)
def create_user(
    payload: UserCreate,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    target_role = payload.role
    # 咨询师只能在其所辖部门新增学员
    if current.role != Role.ADMIN.value and Permission.MANAGE_USERS.value not in (
        current.permissions or []
    ):
        if current.role != Role.COUNSELOR.value or target_role != Role.STUDENT.value:
            raise HTTPException(status_code=403, detail="没有权限新增该类人员")
        dept_ids = descendant_department_ids(db, current.department_id or -1)
        if payload.department_id not in dept_ids:
            raise HTTPException(status_code=403, detail="只能在所辖部门下新增学员")

    if db.query(User).filter(User.username == payload.username).first():
        raise HTTPException(status_code=400, detail="用户名已存在")
    if payload.student_no and db.query(User).filter(
        User.student_no == payload.student_no
    ).first():
        raise HTTPException(status_code=400, detail="学号已存在")
    perms = payload.permissions or DEFAULT_ROLE_PERMISSIONS.get(payload.role, [])
    user = User(
        username=payload.username,
        full_name=payload.full_name,
        student_no=payload.student_no,
        gender=payload.gender,
        birth_date=payload.birth_date,
        role=payload.role,
        permissions=perms,
        department_id=payload.department_id,
        hashed_password=hash_password(payload.password),
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return _to_out(db, user)


@router.put("/{user_id}", response_model=UserOut)
def update_user(
    user_id: int,
    payload: UserUpdate,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    _require_can_manage(db, current, user)
    data = payload.model_dump(exclude_unset=True)
    # 咨询师不能改角色 / 权限，只能维护学员档案信息
    if current.role != Role.ADMIN.value and Permission.MANAGE_USERS.value not in (
        current.permissions or []
    ):
        for k in ("role", "permissions", "is_active"):
            data.pop(k, None)
    if "password" in data and data["password"]:
        user.hashed_password = hash_password(data.pop("password"))
    else:
        data.pop("password", None)
    if data.get("student_no"):
        dup = db.query(User).filter(
            User.student_no == data["student_no"], User.id != user_id
        ).first()
        if dup:
            raise HTTPException(status_code=400, detail="学号已存在")
    if "department_id" in data and data["department_id"] is not None:
        if not db.query(Department).filter(
            Department.id == data["department_id"]
        ).first():
            raise HTTPException(status_code=404, detail="部门不存在")
    for k, v in data.items():
        setattr(user, k, v)
    db.commit()
    db.refresh(user)
    return _to_out(db, user)


@router.post("/{user_id}/move", response_model=UserOut)
def move_user(
    user_id: int,
    department_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(Permission.MANAGE_USERS)),
):
    """将用户挪动到另一个班级 / 部门。"""
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    if not db.query(Department).filter(Department.id == department_id).first():
        raise HTTPException(status_code=404, detail="目标部门不存在")
    user.department_id = department_id
    db.commit()
    db.refresh(user)
    return _to_out(db, user)


@router.put("/{user_id}/permissions", response_model=UserOut)
def set_permissions(
    user_id: int,
    permissions: list[str],
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(Permission.MANAGE_USERS)),
):
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    valid = {p.value for p in Permission}
    user.permissions = [p for p in permissions if p in valid]
    db.commit()
    db.refresh(user)
    return _to_out(db, user)


@router.get("/{user_id}/reports", response_model=list[ReportSummary])
def user_reports(
    user_id: int,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    """查看某位学员的历史测评报告（点击学员信息可见）。"""
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    if not (can_manage_user(db, current, user) or current.id == user_id):
        raise HTTPException(status_code=403, detail="没有权限查看该学员报告")
    reports = (
        db.query(Report)
        .filter(Report.user_id == user_id)
        .order_by(Report.submitted_at.desc())
        .all()
    )
    out = []
    for r in reports:
        # 仅返回当前用户有权查看的报告
        if not can_view_report(db, current, r) and current.id != user_id:
            continue
        out.append(
            ReportSummary(
                id=r.id,
                user_id=r.user_id,
                user_name=(r.user.full_name or r.user.username) if r.user else "",
                student_no=r.user.student_no if r.user else "",
                department_name=r.user.department.name
                if r.user and r.user.department
                else None,
                scale_name=r.scale.name if r.scale else "",
                report_type=r.report_type,
                total_score=r.total_score,
                crisis_level=r.crisis_level,
                counselor_name=(r.counselor.full_name or r.counselor.username)
                if r.counselor
                else "",
                submitted_at=r.submitted_at,
            )
        )
    return out


@router.delete("/{user_id}")
def delete_user(
    user_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(Permission.MANAGE_USERS)),
):
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    user.is_active = False
    db.commit()
    return {"ok": True}


# ---------- 批量导入 ----------
IMPORT_HEADERS = ["学号", "姓名", "性别", "出生日期(YYYY-MM-DD)", "部门/班级名称", "初始密码"]


@router.get("/import/template")
def import_template(
    _: User = Depends(get_current_user),
):
    """下载批量导入学员的 Excel 模板。"""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "学员导入模板"
    ws.append(IMPORT_HEADERS)
    ws.append(["20260101", "张三", "男", "2009-09-01", "高一(1)班", "student123"])
    ws.append(["20260102", "李四", "女", "2009-10-12", "高一(2)班", "student123"])
    for i, _h in enumerate(IMPORT_HEADERS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type=XLSX_MIME,
        headers={
            "Content-Disposition": "attachment; filename=student_import_template.xlsx"
        },
    )


@router.post("/import")
def import_students(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    """批量导入学员：按「部门/班级名称」对应到系统内的分级部门。"""
    is_admin = current.role == Role.ADMIN.value or (
        Permission.MANAGE_USERS.value in (current.permissions or [])
    )
    if not (is_admin or current.role == Role.COUNSELOR.value):
        raise HTTPException(status_code=403, detail="没有权限导入人员")

    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file.file.read()), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="无法解析该文件，请使用提供的 Excel 模板")
    ws = wb.active
    dept_by_name = {d.name: d for d in db.query(Department).all()}
    allowed_dept_ids = None
    if not is_admin:
        allowed_dept_ids = set(
            descendant_department_ids(db, current.department_id or -1)
        )

    created, errors = 0, []
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    for idx, row in enumerate(rows, start=2):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        cells = list(row) + [None] * (6 - len(row))
        student_no, name, gender, birth, dept_name, pwd = cells[:6]
        student_no = str(student_no).strip() if student_no is not None else ""
        if not student_no:
            errors.append(f"第{idx}行：缺少学号")
            continue
        if db.query(User).filter(
            (User.student_no == student_no) | (User.username == student_no)
        ).first():
            errors.append(f"第{idx}行：学号 {student_no} 已存在")
            continue
        dept = dept_by_name.get(str(dept_name).strip()) if dept_name else None
        if dept_name and not dept:
            errors.append(f"第{idx}行：找不到部门/班级「{dept_name}」")
            continue
        if allowed_dept_ids is not None and dept and dept.id not in allowed_dept_ids:
            errors.append(f"第{idx}行：无权导入到「{dept_name}」")
            continue
        gender_val = str(gender).strip() if gender else ""
        if gender_val not in ("男", "女", ""):
            gender_val = ""
        user = User(
            username=student_no,
            student_no=student_no,
            full_name=str(name).strip() if name else "",
            gender=gender_val,
            birth_date=str(birth).strip()[:10] if birth else "",
            role=Role.STUDENT.value,
            permissions=DEFAULT_ROLE_PERMISSIONS[Role.STUDENT.value],
            department_id=dept.id if dept else None,
            hashed_password=hash_password(str(pwd).strip() if pwd else "student123"),
        )
        db.add(user)
        created += 1
    db.commit()
    return {"created": created, "errors": errors}
