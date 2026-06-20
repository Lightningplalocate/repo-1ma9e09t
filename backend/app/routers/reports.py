import io
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from ..auth import get_current_user
from ..database import get_db
from ..models import Report, Scale, TaskAssignment, User
from ..permissions import Permission, Role
from ..schemas import ReportOut, ReportSummary, ReportUpdate, SubmissionIn
from ..services import (
    CRISIS_LABELS,
    build_ai_analysis,
    can_view_report,
    score_submission,
    visible_reports_query,
)

router = APIRouter(prefix="/api/reports", tags=["reports"])

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _counselor_name(r: Report) -> str:
    return (r.counselor.full_name or r.counselor.username) if r.counselor else ""


def _summary(db: Session, r: Report) -> ReportSummary:
    return ReportSummary(
        id=r.id,
        user_id=r.user_id,
        user_name=(r.user.full_name or r.user.username) if r.user else "",
        student_no=r.user.student_no if r.user else "",
        department_name=r.user.department.name if r.user and r.user.department else None,
        scale_name=r.scale.name if r.scale else "",
        report_type=r.report_type,
        total_score=r.total_score,
        crisis_level=r.crisis_level,
        counselor_name=_counselor_name(r),
        submitted_at=r.submitted_at,
    )


def _full_report(db: Session, r: Report) -> ReportOut:
    out = ReportOut.model_validate(r)
    out.user_name = (r.user.full_name or r.user.username) if r.user else ""
    out.student_no = r.user.student_no if r.user else ""
    out.gender = r.user.gender if r.user else ""
    out.birth_date = r.user.birth_date if r.user else ""
    out.department_name = r.user.department.name if r.user and r.user.department else None
    out.scale_name = r.scale.name if r.scale else ""
    out.counselor_name = _counselor_name(r)
    return out


@router.post("/submit", response_model=ReportOut)
def submit(
    payload: SubmissionIn,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    assignment = (
        db.query(TaskAssignment)
        .filter(TaskAssignment.id == payload.assignment_id)
        .first()
    )
    if not assignment:
        raise HTTPException(status_code=404, detail="测评任务不存在")
    if assignment.user_id != current.id:
        raise HTTPException(status_code=403, detail="不能提交他人的测评任务")

    scale = db.query(Scale).filter(Scale.id == assignment.task.scale_id).first()
    answer_map = {a.question_id: a.choice_index for a in payload.answers}
    result = score_submission(scale, answer_map)
    ai_text = build_ai_analysis(current.full_name or current.username, scale, result)

    report = Report(
        assignment_id=assignment.id,
        user_id=current.id,
        scale_id=scale.id,
        report_type="self",
        answers=result["answers"],
        factor_scores=result["factor_scores"],
        total_score=result["total_score"],
        crisis_level=result["crisis_level"],
        ai_analysis=ai_text,
    )
    db.add(report)
    assignment.status = "completed"
    assignment.completed_at = datetime.utcnow()
    db.commit()
    db.refresh(report)
    return _full_report(db, report)


@router.get("", response_model=list[ReportSummary])
def list_reports(
    crisis_only: bool = False,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    """报告管理：按当前用户权限返回可见报告列表。"""
    q = visible_reports_query(db, current)
    if crisis_only:
        q = q.filter(Report.crisis_level.in_(["medium", "high"]))
    reports = q.order_by(Report.submitted_at.desc()).all()
    return [_summary(db, r) for r in reports]


@router.get("/{report_id}", response_model=ReportOut)
def get_report(
    report_id: int,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    report = db.query(Report).filter(Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="报告不存在")
    if not can_view_report(db, current, report):
        raise HTTPException(status_code=403, detail="没有权限查看该报告")
    return _full_report(db, report)


def _can_edit(current: User) -> bool:
    return current.role in (Role.ADMIN.value, Role.COUNSELOR.value) or (
        Permission.EDIT_REPORTS.value in (current.permissions or [])
    )


@router.put("/{report_id}", response_model=ReportOut)
def update_report(
    report_id: int,
    payload: ReportUpdate,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    """咨询师修改报告：咨询师建议、签名咨询师、答案（改答案自动按因子重算结果）。"""
    report = db.query(Report).filter(Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="报告不存在")
    if not _can_edit(current):
        raise HTTPException(status_code=403, detail="没有权限修改报告")
    if not can_view_report(db, current, report):
        raise HTTPException(status_code=403, detail="没有权限修改该报告")

    if payload.counselor_advice is not None:
        report.counselor_advice = payload.counselor_advice
    if payload.counselor_id is not None:
        report.counselor_id = payload.counselor_id or None

    if payload.answers is not None:
        scale = db.query(Scale).filter(Scale.id == report.scale_id).first()
        answer_map = {a.question_id: a.choice_index for a in payload.answers}
        result = score_submission(scale, answer_map)
        report.answers = result["answers"]
        report.factor_scores = result["factor_scores"]
        report.total_score = result["total_score"]
        report.crisis_level = result["crisis_level"]
        report.ai_analysis = build_ai_analysis(
            (report.user.full_name or report.user.username) if report.user else "来访者",
            scale,
            result,
        )
    db.commit()
    db.refresh(report)
    return _full_report(db, report)


def _load_for_export(db: Session, report_id: int, current: User) -> Report:
    report = db.query(Report).filter(Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="报告不存在")
    if not can_view_report(db, current, report):
        raise HTTPException(status_code=403, detail="没有权限导出该报告")
    return report


@router.get("/{report_id}/export/word")
def export_word(
    report_id: int,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    from docx import Document
    from docx.shared import Pt

    r = _load_for_export(db, report_id, current)
    doc = Document()
    doc.add_heading(f"{r.scale.name if r.scale else '心理测评'}报告", level=0)

    doc.add_heading("基本信息", level=1)
    info = [
        ("来访者", (r.user.full_name or r.user.username) if r.user else ""),
        ("学号", r.user.student_no if r.user else ""),
        ("性别", r.user.gender if r.user else ""),
        ("出生日期", r.user.birth_date if r.user else ""),
        ("班级/部门", r.user.department.name if r.user and r.user.department else "-"),
        ("咨询师", _counselor_name(r) or "-"),
        ("总分", str(r.total_score)),
        ("风险等级", CRISIS_LABELS.get(r.crisis_level, r.crisis_level)),
        ("提交时间", r.submitted_at.strftime("%Y-%m-%d %H:%M") if r.submitted_at else ""),
    ]
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, v in info:
        cells = t.add_row().cells
        cells[0].text = k
        cells[1].text = str(v)

    doc.add_heading("因子分析", level=1)
    ft = doc.add_table(rows=1, cols=4)
    ft.style = "Light Grid Accent 1"
    hdr = ft.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "因子", "题目数", "总分", "均分"
    for f in r.factor_scores or []:
        c = ft.add_row().cells
        c[0].text = str(f.get("name", ""))
        c[1].text = str(f.get("items", ""))
        c[2].text = str(f.get("sum", ""))
        c[3].text = str(f.get("score", ""))

    doc.add_heading("AI 分析（第三人称）", level=1)
    doc.add_paragraph(r.ai_analysis or "")

    doc.add_heading("咨询师建议", level=1)
    doc.add_paragraph(r.counselor_advice or "（暂无）")
    sign = doc.add_paragraph()
    run = sign.add_run(f"咨询师签名：{_counselor_name(r) or '____________'}")
    run.font.size = Pt(11)

    doc.add_heading("逐题作答明细", level=1)
    at = doc.add_table(rows=1, cols=5)
    at.style = "Light Grid Accent 1"
    h = at.rows[0].cells
    h[0].text, h[1].text, h[2].text, h[3].text, h[4].text = (
        "题号", "题目", "所属因子", "作答", "得分",
    )
    for a in r.answers or []:
        c = at.add_row().cells
        c[0].text = str(a.get("order", ""))
        c[1].text = str(a.get("text", ""))
        c[2].text = str(a.get("factor_name", ""))
        c[3].text = str(a.get("choice_label", ""))
        c[4].text = str(a.get("score", ""))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    fname = f"report_{report_id}.docx"
    return StreamingResponse(
        buf,
        media_type=DOCX_MIME,
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )


@router.get("/{report_id}/export/excel")
def export_excel(
    report_id: int,
    db: Session = Depends(get_db),
    current: User = Depends(get_current_user),
):
    import openpyxl

    r = _load_for_export(db, report_id, current)
    wb = openpyxl.Workbook()

    # Sheet 1：报告概览（基本信息 + 因子分析 + 咨询师建议）
    ws = wb.active
    ws.title = "报告概览"
    ws.append(["项目", "内容"])
    overview = [
        ("来访者", (r.user.full_name or r.user.username) if r.user else ""),
        ("学号", r.user.student_no if r.user else ""),
        ("性别", r.user.gender if r.user else ""),
        ("出生日期", r.user.birth_date if r.user else ""),
        ("班级/部门", r.user.department.name if r.user and r.user.department else "-"),
        ("量表", r.scale.name if r.scale else ""),
        ("总分", r.total_score),
        ("风险等级", CRISIS_LABELS.get(r.crisis_level, r.crisis_level)),
        ("咨询师", _counselor_name(r)),
        ("咨询师建议", r.counselor_advice or ""),
    ]
    for k, v in overview:
        ws.append([k, v])
    ws.append([])
    ws.append(["因子", "题目数", "总分", "均分"])
    for f in r.factor_scores or []:
        ws.append([f.get("name", ""), f.get("items", ""), f.get("sum", ""), f.get("score", "")])
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 40

    # Sheet 2：逐题作答明细（单独页，表格化）
    ws2 = wb.create_sheet("逐题作答明细")
    ws2.append(["题号", "题目", "所属因子", "作答", "得分"])
    for a in r.answers or []:
        ws2.append([
            a.get("order", ""),
            a.get("text", ""),
            a.get("factor_name", ""),
            a.get("choice_label", ""),
            a.get("score", ""),
        ])
    ws2.column_dimensions["B"].width = 50
    for col in ("A", "C", "D", "E"):
        ws2.column_dimensions[col].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"report_{report_id}.xlsx"
    return StreamingResponse(
        buf,
        media_type=XLSX_MIME,
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )
